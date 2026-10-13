"""Fast tests (no real LLM): limit parsing, output checks, and a full run against a mock API."""

import importlib.util
import os
import subprocess
import sys
import tempfile
import threading
import unittest
from datetime import datetime, timezone
from http.server import HTTPServer
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))  # tomeshift imports its sibling modules (tomeshift_epub, tomeshift_ui)
sys.argv = ["tomeshift"]
spec = importlib.util.spec_from_file_location("tomeshift", ROOT / "tomeshift.py")
ts = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ts)


class LimitParsing(unittest.TestCase):
    def test_detects_limits(self):
        for msg in [
            "Claude AI usage limit reached|1760000000",
            "You've hit your limit · resets 3pm (UTC)",
            "429 RESOURCE_EXHAUSTED",
            "Error: Too Many Requests",
        ]:
            self.assertTrue(ts.LIMIT_RE.search(msg), msg)

    def test_reset_times(self):
        now = datetime.now(timezone.utc)
        self.assertEqual(ts.parse_reset("limit reached|1760000000").year, 2025)
        t = ts.parse_reset("Try again in 2h 5m")
        self.assertAlmostEqual((t - now).total_seconds(), 7500, delta=5)
        self.assertIsNone(ts.parse_reset("please try again later"))  # "again" must not match "in"


class Checks(unittest.TestCase):
    src = "Alice ran.\n\nThe Rabbit was late.\n"

    def test_good_bangla(self):
        out = "অ্যালিস দৌড়াল। " * 4 + "\n\nখরগোশের দেরি হয়ে গেল। " * 3
        self.assertEqual(ts.check(self.src * 3, out, {}, "en", "bn"), [])

    def test_untranslated(self):
        self.assertTrue(ts.check(self.src, self.src, {}, "en", "bn"))

    def test_extra_paragraphs_from_context(self):
        out = "\n\n".join(["অ্যালিস দৌড়াল। খরগোশের দেরি হয়ে গেল।" * 2] * 6)
        self.assertTrue(any("extra paragraphs" in p for p in ts.check(self.src * 2, out, {}, "en", "bn")))

    def test_glossary(self):
        out = "অ্যালিস দৌড়াল। " * 6 + "\n\nদেরি হয়ে গেল। " * 3
        gloss = {"Alice": "এলিস", "Rabbit": "খরগোশ", "late": "দেরি!", "ran": "ছুটল"}
        self.assertTrue(any("glossary" in p for p in ts.check(self.src * 3, out, gloss, "en", "bn")))


class TextPrep(unittest.TestCase):
    def test_clean_dehyphenates_and_rejoins(self):
        paras = ts.clean("A long sen-\ntence that was\n\nsplit across pages.\n\n\n# Chapter 2\n\nNew text.")
        self.assertEqual(paras[0], "A long sentence that was split across pages.")
        self.assertEqual(paras[1], "# Chapter 2")

    def test_tidy_strips_wrappers(self):
        self.assertEqual(ts.tidy("```\nHere is the translation:\n<passage>\nনমস্কার\n</passage>\n```"), "নমস্কার\n")

    def test_lang_info_and_rules(self):
        self.assertEqual(ts.lang_info("bn")[0], "Bangla")
        self.assertEqual(ts.lang_info("Swahili"), ("Swahili", None, ""))
        rules = ts.rules_for("hi", None, "fr")
        self.assertIn("French passage", rules)
        self.assertIn("Hindi", rules)
        self.assertNotIn("{", rules)

    def test_prep_build_epub_roundtrip(self):
        import argparse
        import zipfile

        with tempfile.TemporaryDirectory() as d:
            ts.cmd_prep(
                argparse.Namespace(
                    book=str(ROOT / "samples/alice.txt"), project=d, words=150, lang="bn", **{"from": "en"}
                )
            )
            chunks = sorted(Path(d, "chunks").glob("*.md"))
            self.assertEqual(len(chunks), 3)
            self.assertTrue(chunks[2].read_text().startswith("# CHAPTER II"))
            for c in chunks:  # pretend-translate: copy the source
                Path(d, "out", c.name).write_text(c.read_text())
            ts.cmd_build(argparse.Namespace(project=d, title="Alice"))
            epub = Path(d, "book_bn.epub")
            with zipfile.ZipFile(epub) as z:
                first = z.infolist()[0]
                self.assertEqual((first.filename, first.compress_type), ("mimetype", zipfile.ZIP_STORED))
                self.assertIn("OEBPS/nav.xhtml", z.namelist())
            text = ts.read_epub(epub)  # our own reader can read what we wrote
            self.assertIn("Rabbit-Hole", text)
            self.assertIn("Pool of Tears", text)


def fake_translate(project):
    """Replace every English word with a Bangla word, keeping all markers (what a good model does)."""
    import re

    for c in Path(project, "chunks").glob("*.md"):
        Path(project, "out", c.name).write_text(re.sub(r"[A-Za-z]+", "অনুবাদ", c.read_text()))


class EpubInPlace(unittest.TestCase):
    def setUp(self):
        import argparse

        sys.path.insert(0, str(ROOT / "tests"))
        from epub_fixture import make_epub

        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)
        make_epub(self.d / "alice.epub")
        ts.cmd_prep(argparse.Namespace(book=str(self.d / "alice.epub"), project=str(self.d / "p"), words=400,
                                       lang="bn", flat=False, **{"from": "en"}))  # fmt: skip
        fake_translate(self.d / "p")
        self.build = lambda layout: ts.cmd_build(
            argparse.Namespace(project=str(self.d / "p"), title=None, layout=layout)
        )

    def tearDown(self):
        self.tmp.cleanup()

    def read(self, name, member):
        import zipfile

        with zipfile.ZipFile(self.d / "p" / name) as z:
            return z.read(member)

    def test_segments_and_markers(self):
        chunk = (self.d / "p/chunks/0001.md").read_text()
        self.assertIn("[[2]] Alice was beginning to get ⟦1⟧very⟦/1⟧ tired", chunk)
        self.assertIn("The White Rabbit with a watch", chunk)  # img alt
        self.assertIn("Loose words in a box.", chunk)  # text next to block elements
        self.assertNotIn("do not translate me", chunk)  # <pre><code>
        self.assertNotIn("⟦3⟧1⟦/3⟧", chunk)  # footnote number copied verbatim, not translated

    def test_replace_layout_keeps_everything_else(self):
        import xml.dom.minidom
        import zipfile

        self.build("replace")
        src = self.d / "alice.epub"
        with zipfile.ZipFile(self.d / "p/book_bn.epub") as z, zipfile.ZipFile(src) as orig:
            first = z.infolist()[0]
            self.assertEqual((first.filename, first.compress_type), ("mimetype", zipfile.ZIP_STORED))
            for f in ("OEBPS/style.css", "OEBPS/images/rabbit.png"):
                self.assertEqual(z.read(f), orig.read(f), f)
            for n in z.namelist():
                if n.endswith((".xhtml", ".opf", ".ncx")):
                    xml.dom.minidom.parseString(z.read(n))  # still well-formed
        ch = self.read("book_bn.epub", "OEBPS/text/ch1.xhtml").decode()
        self.assertTrue(ch.startswith('<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>'))
        self.assertIn('lang="bn"', ch)
        self.assertIn("<strong>অনুবাদ <em>অনুবাদ</em></strong>", ch)
        self.assertIn('<a href="#n1" id="r1" epub:type="noteref"><sup>1</sup></a>', ch)
        self.assertIn('print("do not translate me")', ch)
        self.assertIn('alt="অনুবাদ', ch)
        self.assertNotIn("Alice", ch)
        opf = self.read("book_bn.epub", "OEBPS/content.opf").decode()
        self.assertIn("<dc:language>bn</dc:language>", opf)
        self.assertNotIn("Alice in Wonderland", opf)
        self.assertNotIn("Rabbit-Hole", self.read("book_bn.epub", "OEBPS/toc.ncx").decode())
        self.assertNotIn("Rabbit-Hole", self.read("book_bn.epub", "OEBPS/nav.xhtml").decode())

    def test_bilingual_layout(self):
        import re

        self.build("bilingual")
        ch = self.read("book_bn_bilingual.epub", "OEBPS/text/ch1.xhtml").decode()
        self.assertIn("Alice was", ch)  # original kept
        self.assertIn('<p class="first" lang="bn" xml:lang="bn">অনুবাদ', ch)  # translation right after it
        ids = re.findall(r' id="([^"]+)"', ch)
        self.assertEqual(len(ids), len(set(ids)), "duplicate ids")
        md = (self.d / "p/book_bn_bilingual.md").read_text()
        self.assertLess(md.index("Alice was"), md.index("অনুবাদ অনুবাদ অনুবাদ অনুবাদ অনুবাদ"))

    def test_broken_markers_fall_back_safely(self):
        out = self.d / "p/out/0001.md"
        out.write_text(out.read_text().replace("⟦/1⟧", "", 1))  # model dropped a closing marker
        self.build("replace")
        ch = self.read("book_bn.epub", "OEBPS/text/ch1.xhtml").decode()
        self.assertIn('<a href="#n1" id="r1"', ch)  # the footnote link survives the fallback
        self.assertNotIn("⟦", ch)


class Markers(unittest.TestCase):
    def test_segment_and_inline_checks(self):
        src = "[[1]] Hello ⟦1⟧world⟦/1⟧\n\n[[2]] Bye⟦2/⟧"
        self.assertEqual(ts.check_markers(src, "[[1]] হ্যালো ⟦1⟧বিশ্ব⟦/1⟧\n\n[[2]] বিদায়⟦2/⟧"), [])
        # dropping a whole formatting pair is fine (no word to emphasise in the target language)...
        self.assertEqual(ts.check_markers(src, "[[1]] হ্যালো বিশ্ব\n\n[[2]] বিদায়⟦2/⟧"), [])
        # ...but losing an image/note marker, half a pair, or duplicating one is not
        self.assertTrue(ts.check_markers(src, "[[1]] হ্যালো ⟦1⟧বিশ্ব⟦/1⟧\n\n[[2]] বিদায়"))
        self.assertTrue(ts.check_markers(src, "[[1]] হ্যালো ⟦1⟧বিশ্ব\n\n[[2]] বিদায়⟦2/⟧"))
        self.assertTrue(ts.check_markers(src, "[[1]] ⟦1⟧হ্যালো⟦/1⟧ ⟦1⟧বিশ্ব⟦/1⟧\n\n[[2]] বিদায়⟦2/⟧"))
        self.assertTrue(ts.check_markers(src, "[[1]] হ্যালো ⟦1⟧বিশ্ব⟦/1⟧ বিদায়⟦2/⟧"))  # segment 2 merged


class Docx(unittest.TestCase):
    def test_roundtrip(self):
        md = "# Chapter One\n\nAlice was *very* tired.\n\n## Part\n\nঅ্যালিস খুব ক্লান্ত।\n"
        with tempfile.TemporaryDirectory() as d:
            path = Path(d, "b.docx")
            ts.write_docx(md, path, "bn", "Alice")
            self.assertEqual(ts.read_docx(path), md.strip())
            try:
                import docx  # python-docx, if installed, must open it too (= Word-compatible structure)
            except ImportError:
                return
            doc = docx.Document(str(path))
            self.assertEqual(doc.paragraphs[0].style.name, "Heading 1")
            self.assertTrue(doc.paragraphs[1].runs[1].italic)


class EndToEnd(unittest.TestCase):
    def test_mock_api_run_with_429(self):
        sys.path.insert(0, str(ROOT / "tests"))
        import mock_openai

        srv = HTTPServer(("127.0.0.1", 0), mock_openai.H)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        with tempfile.TemporaryDirectory() as d:
            env = {**os.environ, "OPENAI_BASE_URL": f"http://127.0.0.1:{srv.server_port}/v1"}

            def run(*a):
                return subprocess.run(
                    [sys.executable, str(ROOT / "tomeshift.py"), *a],
                    env=env,
                    capture_output=True,
                    text=True,
                    timeout=120,
                )

            r = run("prep", str(ROOT / "samples/alice.txt"), "-p", d, "--words", "150")
            self.assertEqual(r.returncode, 0, r.stderr)
            r = run("run", "-p", d, "--agents", "openai-api", "-j", "2")
            self.assertEqual(r.returncode, 0, r.stderr)
            self.assertIn("paused", r.stdout)  # first request got 429
            self.assertEqual(len(list(Path(d, "out").glob("*.md"))), 3)
            r = run("build", "-p", d)
            self.assertTrue(Path(d, "book_bn.epub").exists())
        srv.shutdown()


class WebUI(unittest.TestCase):
    def test_upload_and_status(self):
        import json
        import time
        import urllib.request

        sys.path.insert(0, str(ROOT))
        import tomeshift_ui

        with tempfile.TemporaryDirectory() as d:
            srv = tomeshift_ui.ThreadingHTTPServer(
                ("127.0.0.1", 0), tomeshift_ui.make_handler(d, {c: v[0] for c, v in ts.LANGS.items()}, ts.BACKENDS)
            )
            threading.Thread(target=srv.serve_forever, daemon=True).start()
            base = f"http://127.0.0.1:{srv.server_port}"

            def get(path):
                with urllib.request.urlopen(base + path) as r:
                    return json.loads(r.read())

            self.assertIn("bn", get("/api/meta")["langs"])
            req = urllib.request.Request(
                base + "/api/prep?p=alice&from=en&lang=bn&words=150",
                method="POST",
                data=(ROOT / "samples/alice.txt").read_bytes(),
                headers={"X-Filename": "alice.txt"},
            )
            urllib.request.urlopen(req)
            for _ in range(60):
                st = get("/api/status?p=alice")
                if st["job"] and st["job"]["ok"] is not None:
                    break
                time.sleep(0.5)
            self.assertTrue(st["job"]["ok"], st["job"])
            self.assertEqual(st["chunks"], 3)
            self.assertEqual(get("/api/projects")[0]["name"], "alice")
            self.assertIn("Source term", urllib.request.urlopen(base + "/").read().decode())
            srv.shutdown()


if __name__ == "__main__":
    unittest.main()
